"""Sprint 4 (C2) gate (fast/mocked slice): extract_from_file() reuses the
shared document reader (no second PDF/DOCX parser) and the existing
extract_from_text() extraction call (no second extraction prompt). The
live end-to-end path (real DOCX + real PDF upload, real extraction, real
save, dedup, drafting integration, critique funder-fit citing) was run
manually against the dev server; see
docs/DASHBOARD-CORRECTIONS-SPRINT4-PASS0.md."""
from __future__ import annotations
import docx

from scriptorium_ingest.funding.crawl_funding import extract_from_file


def test_extract_from_file_parses_docx_via_shared_reader_and_extracts(tmp_path, monkeypatch):
    import scriptorium_ingest.funding.crawl_funding as cf
    monkeypatch.setattr(cf, "complete", lambda prompt, model: (
        '[{"program": "Test Grant", "sponsor": "Test Sponsor", "deadline": "2027-01-01", '
        '"deadline_uncertain": false, "amount": "$1,000", "duration": null, "eligibility": null, '
        '"keywords": [], "priorities": null, "requirements": null, "proposal_structure": null, '
        '"submission_format": null}]'
    ))

    d = docx.Document()
    d.add_heading("Test Grant", level=1)
    d.add_paragraph("Sponsor: Test Sponsor")
    path = tmp_path / "call.docx"
    d.save(str(path))

    result = extract_from_file(str(path))

    assert result["program"] == "Test Grant"
    assert result["sponsor"] == "Test Sponsor"


def test_extract_from_file_falls_back_to_document_title_when_extraction_finds_no_program(tmp_path, monkeypatch):
    """The LLM extraction found nothing usable — rather than an opportunity
    with a blank program name, fall back to the document's own title
    (still not fabricated: it's literally the text at the top of the
    file), same 'degrade gracefully, never fabricate' rule as elsewhere."""
    import scriptorium_ingest.funding.crawl_funding as cf
    monkeypatch.setattr(cf, "complete", lambda prompt, model: "[]")  # extraction found nothing

    d = docx.Document()
    d.add_heading("Untitled Call Document", level=1)
    path = tmp_path / "call.docx"
    d.save(str(path))

    result = extract_from_file(str(path))

    assert result["program"] == "Untitled Call Document"
